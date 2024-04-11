import os
import shutil
import comtypes
from docx import Document
from docx2pdf import convert
from dotenv import dotenv_values


class DocumentProcessor:

    def __init__(self, data):
        self.env_vars = dotenv_values('.env')

        date = '-'.join(data.get('date').split())
        # self.save_directory_path = self.env_vars.get('SAVE_DIRECTORY_PATH')
        self.save_directory_path = "../"

        self.save_directory_path = os.path.join(self.save_directory_path, date)
        
        if not os.path.exists(self.save_directory_path):
           os.makedirs(self.save_directory_path)

        # self.name = self.env_vars.get('NAME')
        self.name = data.get('name')
        self.job_id = data.get('job_id')
        self.type_ = data.get('type')
        self.data = data

        # self.cover_letter_template_path = self.env_vars.get('NON_TECHNICAL_COVER_LETTER_PATH') if self.type_ == 'non_technical' else self.env_vars.get('TECHNICAL_COVER_LETTER_PATH')
        # self.resume_template_path = self.env_vars.get('NON_TECHNICAL_RESUME_PATH') if self.type_ == 'non_technical' else self.env_vars.get('TECHNICAL_RESUME_PATH')

        self.cover_letter_template_path = "./resources/cover-letter/non_technical.docx" if self.type_ == 'non_technical' else "./resources/cover-letter/technical.docx"
        self.resume_template_path = "./resources/resume/non_technical.pdf" if self.type_ == 'non_technical' else "./resources/resume/technical.pdf"
        self.doc = Document(self.cover_letter_template_path)

        placeholders = {
            '[DATE]': 'date',
            '[SENDER\'S ADDRESS]': 'address',
            '[JOB TITLE]': 'job_title',
            '[ORGANIZATION NAME]': 'organization_name',
            '[TEAM NAME]': 'team_name',
            '[DUTIES]': 'duties'
        }

        for placeholder, key in placeholders.items():
            self.process_document(placeholder, data.get(key))

        self.save_document

    def process_document(self, target_word, replacement_word):
        for paragraph in self.doc.paragraphs:
            if target_word in paragraph.text:
                paragraph.text = paragraph.text.replace(target_word, replacement_word)

                if target_word == '[SENDER\'S ADDRESS]':
                    parts = replacement_word.split(', ')
                    formatted_address = ',\n'.join(parts)
                    paragraph.text = formatted_address

            for run in paragraph.runs:
                if target_word in run.text:
                    run.text = run.text.replace(target_word, replacement_word)
        
    def save_document(self):
        directory_name = self.job_id

        parent_directory_path = self.save_directory_path

        save_directory_path = os.path.join(parent_directory_path, directory_name)
        if not os.path.exists(save_directory_path):
            os.makedirs(save_directory_path)

        try:
            name = '_'.join(self.name.split())
            file_name = name + '_' + directory_name
            docx_file_save_path = os.path.join(save_directory_path, file_name + '_Cover_Letter.docx')

            pdf_file_save_path = os.path.join(save_directory_path, file_name  + '_Cover_Letter.pdf')

            resume_destination_path = os.path.join(save_directory_path, file_name + '_Resume.pdf')

            self.doc.save(docx_file_save_path)
        except Exception as e:
            return e

        try:
            shutil.copy(self.resume_template_path, resume_destination_path)
        except Exception as e:
            return e

        if os.path.isfile(pdf_file_save_path):
            os.remove(pdf_file_save_path)

        comtypes.CoInitialize()

        try:
            convert(docx_file_save_path)
        except Exception as e:
            return e
        
        return True
